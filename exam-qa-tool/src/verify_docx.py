"""
Exam Verification Tool — Final Version
Reads Word (.docx) via python-docx, compares against questions.json
Generates ANSWER_VERIFICATION_REPORT.csv and DATA_ERROR_REPORT.md

Usage:
    pip install python-docx openpyxl
    python verify_docx.py
"""

import json, os, re, sys, csv, datetime
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("[ERROR] python-docx not installed.")
    print("  Run: pip install python-docx")
    sys.exit(1)

# ── Paths ──────────────────────────────────────────────────────────────────────
SOURCE_DIR     = Path(r"E:\Đề")
QUESTIONS_JSON = Path(r"E:\Ôn thi\src\data\questions.json")
REPORTS_DIR    = Path(r"E:\Ôn thi\reports")
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

TONG_HOP_NAME  = "TÔNG HỢP TRẮC NGHIỆM.docx"   # actual filename on disk

# ── Load website JSON ──────────────────────────────────────────────────────────
print("[INFO] Đọc questions.json ...")
with open(QUESTIONS_JSON, encoding="utf-8") as f:
    site_exams = json.load(f)
site_map = {e["id"]: e for e in site_exams}

# ── Text normalizer ────────────────────────────────────────────────────────────
def norm(t: str) -> str:
    if not t:
        return ""
    t = t.strip()
    t = re.sub(r"\s+", " ", t)
    for old, new in [("\u2013","-"),("\u2014","-"),("\u2018","'"),("\u2019","'"),
                     ("\u201c",'"'),("\u201d",'"'),("\u00e2","â"),("\u0101","a")]:
        t = t.replace(old, new)
    return t

# ── DOCX parser ────────────────────────────────────────────────────────────────
OPT_RE  = re.compile(r"^([A-Da-d])[.\)]\s*(.*)", re.DOTALL)
Q_RE    = re.compile(r"^(?:Câu\s+)?(\d+)[.:]\s*(.*)", re.DOTALL)
AK_CELL = re.compile(r"(\d+)\s*[.\-:]\s*([ABCD])")   # answer key in table cell

def extract_answer_key_from_tables(doc) -> dict:
    """Try to find answer key in any table in the doc."""
    key = {}
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                ct = cell.text.strip()
                for m in AK_CELL.finditer(ct):
                    key[int(m.group(1))] = m.group(2).upper()
    return key

def extract_answer_key_from_text(paragraphs: list) -> tuple:
    """Scan for a paragraph block that is an answer key. Returns (key_dict, body_end_idx)."""
    key = {}
    body_end = len(paragraphs)
    for i, text in enumerate(paragraphs):
        if re.search(r"(đáp\s*án|ĐÁP\s*ÁN|ANSWER\s*KEY|Bảng\s*đáp)", text, re.I):
            body_end = i
            joined = " ".join(paragraphs[i:])
            for m in re.finditer(r"(\d+)\s*[.\-:]\s*([ABCD])", joined):
                key[int(m.group(1))] = m.group(2).upper()
            break
        # Inline dense answer key like "1.A 2.B 3.C ..."
        dense = re.findall(r"(\d+)\s*\.\s*([ABCD])", text)
        if len(dense) >= 5:
            body_end = min(body_end, i)
            for num, letter in dense:
                key[int(num)] = letter.upper()
    return key, body_end

def parse_docx(docx_path: Path, q_start: int = 1, q_end: int = None) -> list:
    """
    Parse docx and return list of question dicts:
    { num, question, A, B, C, D, correct }
    q_start/q_end: 1-based inclusive range filter (None = no limit)
    """
    doc = Document(str(docx_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs]

    # Answer key: try tables first, then text scan
    answer_key = extract_answer_key_from_tables(doc)
    if not answer_key:
        answer_key, body_end = extract_answer_key_from_text(paragraphs)
        body_paragraphs = paragraphs[:body_end]
    else:
        body_paragraphs = paragraphs

    questions = []
    i = 0
    while i < len(body_paragraphs):
        text = body_paragraphs[i]
        if not text:
            i += 1
            continue

        mq = Q_RE.match(text)
        if not mq:
            i += 1
            continue

        q_num  = int(mq.group(1))
        q_text = mq.group(2).strip()
        i += 1

        # Accumulate multi-line question text
        while i < len(body_paragraphs):
            nxt = body_paragraphs[i]
            if not nxt:
                i += 1
                break
            if OPT_RE.match(nxt) or Q_RE.match(nxt):
                break
            q_text += " " + nxt
            i += 1

        # Collect A/B/C/D options
        opts = {}
        while i < len(body_paragraphs):
            nxt = body_paragraphs[i]
            if not nxt:
                i += 1
                continue
            mo = OPT_RE.match(nxt)
            if mo:
                letter = mo.group(1).upper()
                opts[letter] = mo.group(2).strip()
                i += 1
            else:
                break

        if len(opts) < 2:
            # Sequence/text question — skip, not MCQ
            continue

        q = {
            "num":      q_num,
            "question": norm(q_text),
            "A":        norm(opts.get("A", "")),
            "B":        norm(opts.get("B", "")),
            "C":        norm(opts.get("C", "")),
            "D":        norm(opts.get("D", "")),
            "correct":  answer_key.get(q_num, ""),
        }

        # Filter by range
        if q_end is not None and not (q_start <= q_num <= q_end):
            continue
        if q_end is None and q_num < q_start:
            continue

        questions.append(q)

    return questions

# ── Site question accessor ─────────────────────────────────────────────────────
OPT_PREFIX = re.compile(r"^([ABCD])\.\s*")

def site_q(exam_id: str, idx0: int) -> dict | None:
    exam = site_map.get(exam_id)
    if not exam:
        return None
    qs = exam.get("questions", [])
    if idx0 >= len(qs):
        return None
    q = qs[idx0]
    opts = {}
    for ans in q.get("answers", []):
        m = OPT_PREFIX.match(ans.strip())
        if m:
            opts[m.group(1)] = ans[m.end():].strip()
    return {
        "num":      idx0 + 1,
        "question": norm(q.get("question", "")),
        "type":     q.get("type", "mcq"),
        "A":        norm(opts.get("A", "")),
        "B":        norm(opts.get("B", "")),
        "C":        norm(opts.get("C", "")),
        "D":        norm(opts.get("D", "")),
        "correct":  str(q.get("correct", "")).strip().upper(),
    }

# ── Similarity helper ──────────────────────────────────────────────────────────
def words(s: str) -> set:
    return set(re.findall(r"\w+", s.lower()))

def similarity(a: str, b: str) -> float:
    wa, wb = words(a), words(b)
    if not wa or not wb:
        return 0.0
    return len(wa & wb) / max(len(wa | wb), 1)

# ══════════════════════════════════════════════════════════════════════════════
#  MAIN COMPARISON
# ══════════════════════════════════════════════════════════════════════════════

EXAM_MAP = {
    "02": ("ĐỀ ÔN LUYỆN SỐ 02.docx", 1, None),
    "03": ("ĐỀ ÔN LUYỆN SỐ 03.docx", 1, None),
    "04": ("ĐỀ ÔN LUYỆN SỐ 04.docx", 1, None),
    "05": ("ĐỀ ÔN LUYỆN SỐ 05.docx", 1, None),
    "06": ("ĐỀ ÔN LUYỆN SỐ 06.docx", 1, None),
    "07": ("ĐỀ ÔN LUYỆN SỐ 07.docx", 1, None),
    "08": ("ĐỀ ÔN LUYỆN SỐ 08.docx", 1, None),
    "09": ("ĐỀ ÔN LUYỆN SỐ 09.docx", 1, None),
    "10": ("ĐỀ ÔN LUYỆN SỐ 10.docx", 1, None),
    "11": ("ĐỀ ÔN LUYỆN SỐ 11.docx", 1, None),
    "12": (TONG_HOP_NAME,   1,  70),
    "13": (TONG_HOP_NAME,  71, 140),
    "14": (TONG_HOP_NAME, 141, 210),
    "15": (TONG_HOP_NAME, 211, 280),
}

rows   = []
errors = []
_docx_cache = {}   # path → parsed list

print()
for exam_id, (fname, qs, qe) in EXAM_MAP.items():
    docx_path = SOURCE_DIR / fname
    if not docx_path.exists():
        errors.append(f"❌ [CRITICAL] Đề {exam_id}: File không tồn tại: {fname}")
        print(f"  ✗ Đề {exam_id}: {fname} — NOT FOUND")
        continue

    exam_obj = site_map.get(exam_id)
    if not exam_obj:
        errors.append(f"❌ [CRITICAL] Đề {exam_id}: Không có trong questions.json")
        print(f"  ✗ Đề {exam_id}: missing from questions.json")
        continue

    # Parse (or reuse cache)
    cache_key = str(docx_path)
    if cache_key not in _docx_cache:
        print(f"  [PARSE] {fname} ...")
        try:
            _docx_cache[cache_key] = parse_docx(docx_path)
        except Exception as exc:
            errors.append(f"❌ [PARSE ERROR] Đề {exam_id}: {exc}")
            print(f"  ✗ Parse error: {exc}")
            continue

    all_word_qs = _docx_cache[cache_key]
    if qe:
        word_qs = [q for q in all_word_qs if qs <= q["num"] <= qe]
        # Re-number relative to 1 for Đề 13,14,15
        offset = qs - 1
        word_qs_norm = [{**q, "num": q["num"] - offset} for q in word_qs]
    else:
        word_qs_norm = all_word_qs

    site_total = len(exam_obj.get("questions", []))
    word_total = len(word_qs_norm)

    print(f"  Đề {exam_id}: Word={word_total} câu MCQ | Site={site_total} câu tổng")

    if word_total == 0:
        errors.append(f"⚠️ [WARN] Đề {exam_id}: Không parse được câu nào từ Word (cần kiểm tra thủ công định dạng DOCX)")

    if word_total != site_total:
        errors.append(
            f"⚠️ [COUNT_MISMATCH] Đề {exam_id}: Word MCQ={word_total}, Site total={site_total} "
            f"(Site có thể chứa thêm câu sequence/text)"
        )

    # Compare each MCQ question
    for wq in word_qs_norm:
        q_num = wq["num"]
        sq    = site_q(exam_id, q_num - 1)   # 0-based

        if sq is None:
            errors.append(f"❌ [MISSING_IN_SITE] Đề {exam_id} Câu {q_num}: Có trong Word nhưng không có trong questions.json")
            rows.append({"Đề": exam_id, "Câu": q_num,
                         "Nội dung câu hỏi (Word)": wq["question"][:100],
                         "Đáp án đúng (Word)": wq["correct"] or "(no key)",
                         "Đáp án đúng (Website)": "N/A",
                         "Trạng thái": "MISSING_IN_SITE"})
            continue

        # Skip non-MCQ site questions when comparing
        if sq["type"] != "mcq":
            continue

        # Answer comparison
        w_ans  = wq["correct"].upper() if wq["correct"] else ""
        s_ans  = sq["correct"].upper() if sq["correct"] else ""

        if not w_ans:
            status = "NO_KEY_IN_WORD"
        elif w_ans == s_ans:
            status = "PASS"
        else:
            status = "ANSWER_MISMATCH"
            errors.append(
                f"❌ [ANSWER_WRONG] Đề {exam_id} Câu {q_num}: "
                f"Word={w_ans} | Website={s_ans}"
            )

        # Question text similarity
        sim = similarity(wq["question"], sq["question"])
        if sim < 0.5 and wq["question"] and sq["question"]:
            errors.append(
                f"⚠️ [TEXT_MISMATCH] Đề {exam_id} Câu {q_num} (similarity={sim:.0%}):\n"
                f"    Word:    {wq['question'][:120]}\n"
                f"    Website: {sq['question'][:120]}"
            )
            if status == "PASS":
                status = "TEXT_DIFFERS"

        rows.append({
            "Đề":                      exam_id,
            "Câu":                     q_num,
            "Nội dung câu hỏi (Word)": wq["question"][:120],
            "Đáp án đúng (Word)":      w_ans or "(no key)",
            "Đáp án đúng (Website)":   s_ans or "(empty)",
            "Trạng thái":              status,
        })

# ══════════════════════════════════════════════════════════════════════════════
#  SAVE REPORTS
# ══════════════════════════════════════════════════════════════════════════════

# CSV
csv_path = REPORTS_DIR / "ANSWER_VERIFICATION_REPORT.csv"
fieldnames = ["Đề","Câu","Nội dung câu hỏi (Word)","Đáp án đúng (Word)","Đáp án đúng (Website)","Trạng thái"]
with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=fieldnames)
    w.writeheader()
    w.writerows(rows)
print(f"\n[SAVED] {csv_path}")

# Optional: XLSX
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Answer Verification"
    ws.append(fieldnames)
    # Style header
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    STATUS_COLORS = {
        "PASS":             "C6EFCE",
        "ANSWER_MISMATCH":  "FFC7CE",
        "TEXT_DIFFERS":     "FFEB9C",
        "NO_KEY_IN_WORD":   "D9D9D9",
        "MISSING_IN_SITE":  "FF0000",
        "MISSING_IN_WORD":  "FF0000",
    }
    for row in rows:
        ws.append([row[f] for f in fieldnames])
        color = STATUS_COLORS.get(row["Trạng thái"], "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)
    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 20
    xlsx_path = REPORTS_DIR / "ANSWER_VERIFICATION_REPORT.xlsx"
    wb.save(xlsx_path)
    print(f"[SAVED] {xlsx_path}")
except ImportError:
    print("[INFO] openpyxl không có — chỉ lưu CSV (pip install openpyxl để có XLSX)")

# Markdown error report
now   = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
total = len(rows)
passed    = sum(1 for r in rows if r["Trạng thái"] == "PASS")
no_key    = sum(1 for r in rows if r["Trạng thái"] == "NO_KEY_IN_WORD")
ans_err   = sum(1 for r in rows if r["Trạng thái"] == "ANSWER_MISMATCH")
text_warn = sum(1 for r in rows if r["Trạng thái"] == "TEXT_DIFFERS")
critical  = [e for e in errors if "[CRITICAL]" in e]
warns     = [e for e in errors if e not in critical]

md = f"""# BÁO CÁO LỖI DỮ LIỆU — DATA ERROR REPORT
> **Ngày tạo**: {now}
> **Nguồn Word**: `E:\\Đề\\*.docx`
> **Nguồn JSON**: `E:\\Ôn thi\\src\\data\\questions.json`

---

## 📊 Tổng Quan

| Chỉ số | Giá trị |
|---|---|
| Tổng câu MCQ so sánh | **{total}** |
| ✅ PASS (khớp hoàn toàn) | **{passed}** |
| ❌ ANSWER_MISMATCH (sai đáp án) | **{ans_err}** |
| ⚠️ TEXT_DIFFERS (nội dung khác) | **{text_warn}** |
| 🔘 NO_KEY_IN_WORD (Word thiếu đáp án) | **{no_key}** |
| ❌ CRITICAL (file/ID không tồn tại) | **{len(critical)}** |

---

## ❌ Lỗi Nghiêm Trọng (CRITICAL)

"""
if critical:
    for e in critical:
        md += f"- {e}\n"
else:
    md += "_Không có lỗi nghiêm trọng._\n"

md += "\n---\n\n## ⚠️ Chi Tiết Các Vấn Đề Phát Hiện\n\n"
if warns:
    for e in warns:
        md += f"{e}\n\n"
else:
    md += "_Không có lỗi bổ sung._\n"

md += """
---

## 📋 Hướng Dẫn Xử Lý

| Trạng thái | Ý nghĩa | Hành động |
|---|---|---|
| `PASS` | Đáp án khớp hoàn toàn | Không cần làm gì |
| `ANSWER_MISMATCH` | Đáp án trong JSON sai so với Word | **Cần xác nhận và sửa** |
| `TEXT_DIFFERS` | Nội dung câu hỏi khác >50% | Kiểm tra thủ công |
| `NO_KEY_IN_WORD` | File Word không có bảng đáp án | Kiểm tra thủ công file Word |
| `MISSING_IN_SITE` | Câu có trong Word nhưng thiếu trên website | Cần thêm vào JSON |

> ⚠️ Mọi sửa đổi cần được phê duyệt thủ công. Tool chỉ **phát hiện và báo cáo**.
"""

md_path = REPORTS_DIR / "DATA_ERROR_REPORT.md"
with open(md_path, "w", encoding="utf-8") as f:
    f.write(md)
print(f"[SAVED] {md_path}")

print(f"""
╔══════════════════════════════════════╗
║  HOÀN THÀNH                          ║
╠══════════════════════════════════════╣
║  Tổng câu:     {total:<6}               ║
║  PASS:         {passed:<6}               ║
║  Lỗi đáp án:  {ans_err:<6}               ║
║  Text khác:    {text_warn:<6}               ║
║  Critical:     {len(critical):<6}               ║
╚══════════════════════════════════════╝
""")

